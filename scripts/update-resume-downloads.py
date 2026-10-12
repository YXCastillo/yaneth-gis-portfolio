from pathlib import Path
import subprocess
from docx import Document

root = Path(__file__).resolve().parents[1]
folder = root / "public" / "resume"
docx_path = folder / "Yaneth-Castillo-GISP-Resume.docx"
pdf_path = folder / "Yaneth-Castillo-GISP-Resume.pdf"
doc = Document(docx_path)

replacements = [
    ("Develop GIS deliverables and public-facing mapping", "Deliver GIS solutions across transportation, utilities, land development, municipal planning, and regulatory projects, translating complex spatial requirements into reliable datasets, interactive applications, and client-facing deliverables."),
    ("Reconstructed a 224.109-acre", "Reconstruct and validate complex legal boundaries from survey documentation, including a 224.109-acre tract defined by 109 metes-and-bounds courses, identifying a 0.65-acre discrepancy in reference data."),
    ("Adapted and extended JSON configurations", "Configure and enhance a public-facing Web GIS platform supporting 133 spatial layers, maintaining separate aerial and street-map configurations with structured layer hierarchies, interactive legends, and visibility controls."),
    ("Build and maintain ArcGIS Online", "Develop and maintain ArcGIS Online and Experience Builder applications, incorporating hosted feature services, vector tile layers, Arcade expressions, and access controls to support internal and public-facing workflows."),
    ("Troubleshoot GIS data integration", "Integrate and troubleshoot municipal infrastructure datasets within enterprise GIS environments, resolving schema incompatibilities, field mapping requirements, coded-value domains, and versioned geodatabase constraints."),
]

def all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs

found = set()
for p in all_paragraphs(doc):
    original = p.text.strip()
    if not original:
        continue
    if "FEB 2024" in original.upper() or "FEBRUARY 2024" in original.upper():
        for run in p.runs:
            run.text = run.text.replace("FEB 2024", "JAN 2024").replace("Feb 2024", "Jan 2024").replace("February 2024", "January 2024").replace("FEBRUARY 2024", "JANUARY 2024")
    for i, (prefix, replacement) in enumerate(replacements):
        if prefix.casefold() in original.casefold():
            # Preserve paragraph style and the formatting of the initial run.
            if p.runs:
                p.runs[0].text = replacement
                for run in p.runs[1:]:
                    run.text = ""
            else:
                p.text = replacement
            found.add(i)
            break

if len(found) != len(replacements):
    raise RuntimeError(f"Expected 5 LJA bullets; found {sorted(found)}. Original files unchanged.")
doc.save(docx_path)
subprocess.run(["libreoffice", "-env:UserInstallation=file:///tmp/lo-resume-ci", "--headless", "--convert-to", "pdf", "--outdir", str(folder), str(docx_path)], check=True)
if not pdf_path.exists() or pdf_path.stat().st_size < 1000:
    raise RuntimeError("PDF conversion failed")
print(f"Updated DOCX and PDF: {docx_path}, {pdf_path}")
